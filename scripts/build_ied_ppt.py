"""Build a 10-slide PPT on Income Elasticity of Demand."""
import os
import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

OUT_DIR = os.path.dirname(os.path.abspath(__file__))
IMG_DIR = os.path.join(OUT_DIR, "_ied_imgs")
os.makedirs(IMG_DIR, exist_ok=True)
PPT_PATH = os.path.join(os.path.dirname(OUT_DIR), "Income_Elasticity_of_Demand.pptx")

# Brand palette
NAVY = RGBColor(0x0F, 0x2C, 0x4A)
TEAL = RGBColor(0x10, 0x9E, 0x8F)
GOLD = RGBColor(0xE8, 0xB5, 0x3B)
LIGHT = RGBColor(0xF4, 0xF7, 0xFA)
DARK = RGBColor(0x1F, 0x2A, 0x3A)
GRAY = RGBColor(0x55, 0x66, 0x77)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# Matplotlib styling
plt.rcParams.update({
    "font.family": "DejaVu Sans",
    "axes.edgecolor": "#1F2A3A",
    "axes.labelcolor": "#1F2A3A",
    "xtick.color": "#1F2A3A",
    "ytick.color": "#1F2A3A",
    "axes.titleweight": "bold",
})

def make_graph(filename, kind):
    """Generate a small clean graph for each elasticity type."""
    fig, ax = plt.subplots(figsize=(4.2, 3.0), dpi=180)
    income = np.linspace(1, 10, 100)
    if kind == "zero":
        q = np.full_like(income, 5.0)
        color = "#6B7280"
        title = "Zero (IED = 0)"
    elif kind == "negative":
        q = 8 - 0.6 * income
        color = "#D9534F"
        title = "Negative (IED < 0)"
    elif kind == "unit":
        q = income  # proportional
        color = "#109E8F"
        title = "Unit (IED = 1)"
    elif kind == "less":
        q = 2 + 1.5 * np.log(income)
        color = "#0F2C4A"
        title = "Necessity (0 < IED < 1)"
    elif kind == "greater":
        q = 0.3 * income ** 1.6
        color = "#E8B53B"
        title = "Luxury (IED > 1)"
    else:
        q = income
        color = "#000"
        title = ""
    ax.plot(income, q, color=color, linewidth=3)
    ax.set_xlabel("Income", fontsize=10)
    ax.set_ylabel("Quantity Demanded", fontsize=10)
    ax.set_title(title, fontsize=11, color="#0F2C4A")
    ax.set_xticks([])
    ax.set_yticks([])
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.grid(True, linestyle="--", alpha=0.25)
    fig.tight_layout()
    path = os.path.join(IMG_DIR, filename)
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path

def make_elasticity_overview(filename):
    fig, ax = plt.subplots(figsize=(5.5, 3.2), dpi=180)
    x = np.linspace(1, 10, 100)
    ax.plot(x, 10 - x, color="#0F2C4A", linewidth=3, label="Demand")
    ax.set_xlabel("Price / Income", fontsize=10)
    ax.set_ylabel("Quantity Demanded", fontsize=10)
    ax.set_title("Elasticity: How Demand Responds to Change", fontsize=11, color="#0F2C4A")
    ax.set_xticks([]); ax.set_yticks([])
    ax.spines["top"].set_visible(False); ax.spines["right"].set_visible(False)
    ax.grid(True, linestyle="--", alpha=0.25)
    fig.tight_layout()
    path = os.path.join(IMG_DIR, filename)
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path

# Pre-generate images
img_overview = make_elasticity_overview("overview.png")
img_zero = make_graph("zero.png", "zero")
img_neg = make_graph("negative.png", "negative")
img_unit = make_graph("unit.png", "unit")
img_less = make_graph("less.png", "less")
img_greater = make_graph("greater.png", "greater")

# Build presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]


def add_bg(slide, color=LIGHT):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    bg.line.fill.background()
    bg.fill.solid(); bg.fill.fore_color.rgb = color
    bg.shadow.inherit = False
    return bg


def add_accent_bar(slide, color=TEAL):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.35), SH)
    bar.line.fill.background()
    bar.fill.solid(); bar.fill.fore_color.rgb = color


def add_text(slide, left, top, width, height, text, *, size=18, bold=False,
             color=DARK, align=PP_ALIGN.LEFT, font="Calibri"):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    lines = text if isinstance(text, list) else [text]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return tb


def add_bullets(slide, left, top, width, height, items, *, size=18, color=DARK,
                bullet_color=TEAL):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(8)
        # Bullet glyph run
        r1 = p.add_run()
        r1.text = "▸  "
        r1.font.name = "Calibri"
        r1.font.size = Pt(size)
        r1.font.bold = True
        r1.font.color.rgb = bullet_color
        r2 = p.add_run()
        r2.text = item
        r2.font.name = "Calibri"
        r2.font.size = Pt(size)
        r2.font.color.rgb = color
    return tb


def add_section_header(slide, eyebrow, title):
    add_text(slide, Inches(0.7), Inches(0.45), Inches(10), Inches(0.4),
             eyebrow.upper(), size=14, bold=True, color=TEAL)
    add_text(slide, Inches(0.7), Inches(0.85), Inches(12), Inches(0.9),
             title, size=34, bold=True, color=NAVY)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.75),
                                  Inches(1.2), Inches(0.06))
    line.line.fill.background(); line.fill.solid(); line.fill.fore_color.rgb = GOLD


def add_footer(slide, page_num, total=10):
    add_text(slide, Inches(0.7), Inches(7.05), Inches(8), Inches(0.3),
             "Income Elasticity of Demand  •  Business / Economics",
             size=10, color=GRAY)
    add_text(slide, Inches(11.8), Inches(7.05), Inches(1.2), Inches(0.3),
             f"{page_num} / {total}", size=10, color=GRAY, align=PP_ALIGN.RIGHT)


# ---------- Slide 1: Title ----------
s = prs.slides.add_slide(BLANK)
add_bg(s, NAVY)
# decorative shapes
circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.5), Inches(-1.5), Inches(5), Inches(5))
circ.line.fill.background(); circ.fill.solid(); circ.fill.fore_color.rgb = TEAL
circ2 = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(-1.5), Inches(5.5), Inches(3.5), Inches(3.5))
circ2.line.fill.background(); circ2.fill.solid(); circ2.fill.fore_color.rgb = GOLD

add_text(s, Inches(0.9), Inches(2.0), Inches(10), Inches(0.5),
         "BUSINESS  •  ECONOMICS  •  UNIVERSITY LEVEL",
         size=16, bold=True, color=GOLD)
add_text(s, Inches(0.9), Inches(2.6), Inches(11.5), Inches(1.6),
         "Income Elasticity of Demand",
         size=54, bold=True, color=WHITE)
add_text(s, Inches(0.9), Inches(4.3), Inches(11), Inches(0.8),
         "How consumer demand responds to changes in income",
         size=22, color=LIGHT)
bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.9), Inches(5.3),
                         Inches(1.5), Inches(0.08))
bar.line.fill.background(); bar.fill.solid(); bar.fill.fore_color.rgb = GOLD
add_text(s, Inches(0.9), Inches(5.5), Inches(10), Inches(0.5),
         "A 10-slide overview", size=16, color=LIGHT)

# ---------- Slide 2: What is Elasticity? ----------
s = prs.slides.add_slide(BLANK)
add_bg(s); add_accent_bar(s)
add_section_header(s, "Foundations", "What is Elasticity?")
add_bullets(s, Inches(0.7), Inches(2.1), Inches(6.5), Inches(4.5), [
    "Measures responsiveness of one variable to a change in another",
    "Common in economics: demand and supply react to price, income, etc.",
    "Expressed as a ratio of % changes — a pure number, no units",
    "Helps firms set prices and governments design policy",
    "Three key types: Price, Cross, and Income elasticity",
], size=20)
s.shapes.add_picture(img_overview, Inches(7.6), Inches(2.2), width=Inches(5.2))
add_footer(s, 2)

# ---------- Slide 3: Introducing IED ----------
s = prs.slides.add_slide(BLANK)
add_bg(s); add_accent_bar(s)
add_section_header(s, "Concept", "Income Elasticity of Demand (IED)")
add_bullets(s, Inches(0.7), Inches(2.1), Inches(12), Inches(4.5), [
    "Definition: % change in quantity demanded ÷ % change in consumer income",
    "Tells us how sensitive demand is when buyers' income rises or falls",
    "Sign of IED reveals the type of good (normal, inferior, etc.)",
    "Size of IED reveals the strength of the response",
    "Examples: smartphones, restaurant meals, public transport tickets",
], size=20)
add_footer(s, 3)

# ---------- Slide 4: Formula + Example ----------
s = prs.slides.add_slide(BLANK)
add_bg(s); add_accent_bar(s)
add_section_header(s, "Formula", "How to Calculate IED")

# Formula card
card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(2.1),
                          Inches(12), Inches(1.6))
card.adjustments[0] = 0.1
card.line.fill.background(); card.fill.solid(); card.fill.fore_color.rgb = NAVY
add_text(s, Inches(0.9), Inches(2.25), Inches(11.6), Inches(0.5),
         "FORMULA", size=14, bold=True, color=GOLD)
add_text(s, Inches(0.9), Inches(2.65), Inches(11.6), Inches(1.0),
         "IED  =   % change in Quantity Demanded   ÷   % change in Income",
         size=24, bold=True, color=WHITE)

add_text(s, Inches(0.7), Inches(4.0), Inches(12), Inches(0.5),
         "Quick example", size=18, bold=True, color=NAVY)
add_bullets(s, Inches(0.7), Inches(4.4), Inches(12), Inches(2.5), [
    "Income rises by 10%; demand for dining out rises by 20%",
    "IED = 20% ÷ 10% = 2.0",
    "Result > 1 → dining out is a luxury good (income-elastic)",
], size=20)
add_footer(s, 4)

# ---------- Helper for Type slides ----------
def type_slide(page, eyebrow, title, definition, example, img_path, accent=TEAL):
    s = prs.slides.add_slide(BLANK)
    add_bg(s); add_accent_bar(s, accent)
    add_text(s, Inches(0.7), Inches(0.45), Inches(10), Inches(0.4),
             eyebrow.upper(), size=14, bold=True, color=accent)
    add_text(s, Inches(0.7), Inches(0.85), Inches(12), Inches(0.9),
             title, size=32, bold=True, color=NAVY)
    line = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.75),
                              Inches(1.2), Inches(0.06))
    line.line.fill.background(); line.fill.solid(); line.fill.fore_color.rgb = GOLD

    # Left bullets
    add_text(s, Inches(0.7), Inches(2.0), Inches(7), Inches(0.5),
             "Explanation", size=16, bold=True, color=accent)
    add_bullets(s, Inches(0.7), Inches(2.4), Inches(7), Inches(2.0),
                definition, size=18, bullet_color=accent)
    add_text(s, Inches(0.7), Inches(4.6), Inches(7), Inches(0.5),
             "Real-life example", size=16, bold=True, color=accent)
    add_bullets(s, Inches(0.7), Inches(5.0), Inches(7), Inches(2.0),
                example, size=18, bullet_color=accent)

    # Right graph
    s.shapes.add_picture(img_path, Inches(8.0), Inches(2.2), width=Inches(4.9))
    add_footer(s, page)
    return s

# ---------- Slide 5: Zero IED ----------
type_slide(5, "Type 1", "Zero Income Elasticity (IED = 0)",
           [
               "Quantity demanded does NOT change when income changes",
               "Demand is completely income-inelastic",
               "Typically basic items needed in fixed amounts",
           ],
           [
               "Salt: families buy the same amount whether rich or poor",
               "Matchsticks, basic toothpaste",
           ],
           img_zero, accent=GRAY)

# ---------- Slide 6: Negative IED ----------
type_slide(6, "Type 2", "Negative Income Elasticity (IED < 0)",
           [
               "As income rises, quantity demanded FALLS",
               "Indicates an inferior good",
               "Buyers switch to better alternatives when richer",
           ],
           [
               "Public bus travel — replaced by cars or cabs",
               "Instant noodles, used clothing",
           ],
           img_neg, accent=RGBColor(0xD9, 0x53, 0x4F))

# ---------- Slide 7: Unit IED ----------
type_slide(7, "Type 3", "Unit Income Elasticity (IED = 1)",
           [
               "Demand changes EXACTLY in proportion to income",
               "10% rise in income → 10% rise in quantity demanded",
               "Spending share on the good stays the same",
           ],
           [
               "Mid-range clothing for many households",
               "Standard household groceries in some markets",
           ],
           img_unit, accent=TEAL)

# ---------- Slide 8: 0 < IED < 1 (Necessities) ----------
type_slide(8, "Type 4", "Less Than One: Necessities (0 < IED < 1)",
           [
               "Demand rises with income, but more SLOWLY",
               "Income-inelastic — people already buy enough",
               "Typical for basic needs",
           ],
           [
               "Food staples like rice, wheat, milk",
               "Electricity and water for the home",
           ],
           img_less, accent=NAVY)

# ---------- Slide 9: IED > 1 (Luxury) ----------
type_slide(9, "Type 5", "Greater Than One: Luxury Goods (IED > 1)",
           [
               "Demand rises FASTER than income",
               "Income-elastic — highly sensitive to income changes",
               "Demand falls sharply in a recession",
           ],
           [
               "Foreign holidays, fine dining, designer brands",
               "Premium cars, jewellery, smartphones",
           ],
           img_greater, accent=GOLD)

# ---------- Slide 10: Conclusion + Thank You ----------
s = prs.slides.add_slide(BLANK)
add_bg(s, NAVY)
circ = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(-2), Inches(-2), Inches(5), Inches(5))
circ.line.fill.background(); circ.fill.solid(); circ.fill.fore_color.rgb = TEAL
circ2 = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(11), Inches(5), Inches(4), Inches(4))
circ2.line.fill.background(); circ2.fill.solid(); circ2.fill.fore_color.rgb = GOLD

add_text(s, Inches(0.9), Inches(0.7), Inches(10), Inches(0.5),
         "WRAP-UP", size=14, bold=True, color=GOLD)
add_text(s, Inches(0.9), Inches(1.1), Inches(11.5), Inches(0.9),
         "Key Takeaways", size=38, bold=True, color=WHITE)

add_bullets(s, Inches(0.9), Inches(2.2), Inches(11.5), Inches(3.5), [
    "IED links income changes to demand changes",
    "Sign tells the type: negative = inferior, positive = normal",
    "Size tells the strength: <1 necessity, =1 unit, >1 luxury",
    "Useful for firms, marketers, and policy makers",
], size=22, color=WHITE, bullet_color=GOLD)

add_text(s, Inches(0.9), Inches(5.9), Inches(11.5), Inches(1.0),
         "Thank You",
         size=44, bold=True, color=GOLD)
add_text(s, Inches(0.9), Inches(6.7), Inches(11.5), Inches(0.5),
         "Questions & Discussion",
         size=18, color=LIGHT)

prs.save(PPT_PATH)
print("Saved:", PPT_PATH)
